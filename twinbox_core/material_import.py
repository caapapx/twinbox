"""Optional spreadsheet/doc → pack fragment importer."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def import_material(path: Path, *, intent: str = "reference") -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            return {"ok": False, "error": "openpyxl is not installed", "recovery": "pip install openpyxl"}
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(max_row=50, values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
        fragment = {"attention_hints": [{"id": path.stem, "utterances": rows[:20], "intent": intent}]}
        return {"ok": True, "fragment": fragment, "source": str(path)}
    if suffix in {".docx"}:
        try:
            import docx  # type: ignore
        except ImportError:
            return {"ok": False, "error": "python-docx is not installed", "recovery": "pip install python-docx"}
        document = docx.Document(str(path))
        paras = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        fragment = {"attention_hints": [{"id": path.stem, "utterances": paras[:20], "intent": intent}]}
        return {"ok": True, "fragment": fragment, "source": str(path)}
    text = path.read_text(encoding="utf-8", errors="replace")[:4000]
    return {
        "ok": True,
        "fragment": {"attention_hints": [{"id": path.stem, "utterances": [text[:200]], "intent": intent}]},
        "source": str(path),
    }
